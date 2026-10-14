"""Stream worksheet rows from an ``.xlsx`` file into Arrow record batches.

openpyxl's read-only mode parses the sheet XML lazily, so memory is bounded by
``batch_size`` rather than by the size of the worksheet.
"""

from __future__ import annotations

import logging
from collections.abc import Iterator, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
import pyarrow as pa

from purchase_patterns.ingestion.schema import (
    COLUMN_CONVERTERS,
    COLUMN_MAP,
    RAW_SCHEMA,
    SOURCE_HEADER,
    CellConversionError,
    IngestionError,
)

logger = logging.getLogger(__name__)

DEFAULT_BATCH_SIZE = 50_000
FIRST_DATA_ROW = 2

_SOURCE_COLUMNS: tuple[str, ...] = tuple(COLUMN_MAP[name] for name in SOURCE_HEADER)
_CONVERTERS = tuple(COLUMN_CONVERTERS[name] for name in _SOURCE_COLUMNS)
_N_SOURCE_COLUMNS = len(SOURCE_HEADER)


@dataclass
class SheetStats:
    """Running counters collected while a worksheet is streamed."""

    sheet_name: str
    header: list[Any] = field(default_factory=list)
    declared_max_row: int | None = None
    declared_max_column: int | None = None
    data_rows_read: int = 0
    blank_rows_skipped: int = 0
    batches: int = 0


def list_sheet_names(workbook_path: Path) -> list[str]:
    """Return the worksheet names of ``workbook_path`` without loading any cells."""
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _normalise_header(row: Sequence[Any]) -> list[Any]:
    header = list(row)
    while header and header[-1] is None:
        header.pop()
    return header


def _fit_row(row: Sequence[Any], sheet_name: str, row_number: int) -> Sequence[Any]:
    if len(row) == _N_SOURCE_COLUMNS:
        return row
    if len(row) < _N_SOURCE_COLUMNS:
        return tuple(row) + (None,) * (_N_SOURCE_COLUMNS - len(row))
    extra = [value for value in row[_N_SOURCE_COLUMNS:] if value is not None]
    if extra:
        raise IngestionError(
            f"Sheet {sheet_name!r} row {row_number}: unexpected values beyond "
            f"column {_N_SOURCE_COLUMNS}: {extra!r}"
        )
    return row[:_N_SOURCE_COLUMNS]


def _build_batch(
    columns: list[list[Any]], sheet_name: str, row_numbers: list[int]
) -> pa.RecordBatch:
    arrays = [
        pa.array(values, type=RAW_SCHEMA.field(name).type)
        for name, values in zip(_SOURCE_COLUMNS, columns, strict=True)
    ]
    arrays.append(pa.array([sheet_name] * len(row_numbers), type=pa.string()))
    arrays.append(pa.array(row_numbers, type=pa.int32()))
    return pa.RecordBatch.from_arrays(arrays, schema=RAW_SCHEMA)


def iter_sheet_batches(
    workbook_path: Path,
    sheet_name: str,
    *,
    batch_size: int = DEFAULT_BATCH_SIZE,
    stats: SheetStats | None = None,
) -> Iterator[pa.RecordBatch]:
    """Yield ``RAW_SCHEMA`` record batches for one worksheet.

    The first row must match ``SOURCE_HEADER`` exactly. Fully blank rows are
    skipped and counted in ``stats``; every other row is converted cell by cell
    and tagged with its sheet name and 1-based Excel row number.

    Raises:
        IngestionError: if the sheet is missing, the header is unexpected, or a
            cell cannot be converted to its declared type.
    """
    if batch_size < 1:
        raise ValueError(f"batch_size must be positive, got {batch_size}")
    stats = stats if stats is not None else SheetStats(sheet_name=sheet_name)

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise IngestionError(
                f"Sheet {sheet_name!r} not found in {workbook_path.name}; "
                f"available sheets: {workbook.sheetnames}"
            )
        worksheet = workbook[sheet_name]
        stats.declared_max_row = worksheet.max_row
        stats.declared_max_column = worksheet.max_column

        rows = worksheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise IngestionError(f"Sheet {sheet_name!r} is empty; expected a header row")
        stats.header = _normalise_header(header_row)
        if tuple(stats.header) != SOURCE_HEADER:
            raise IngestionError(
                f"Sheet {sheet_name!r} header {stats.header!r} does not match "
                f"expected {list(SOURCE_HEADER)!r}"
            )

        columns: list[list[Any]] = [[] for _ in _SOURCE_COLUMNS]
        row_numbers: list[int] = []
        for row_number, raw_row in enumerate(rows, start=FIRST_DATA_ROW):
            row = _fit_row(raw_row, sheet_name, row_number)
            if all(value is None for value in row):
                stats.blank_rows_skipped += 1
                continue

            for index, value in enumerate(row):
                try:
                    columns[index].append(_CONVERTERS[index](value))
                except CellConversionError as exc:
                    raise IngestionError(
                        f"Sheet {sheet_name!r} row {row_number} column "
                        f"{SOURCE_HEADER[index]!r}: {exc}"
                    ) from exc
            row_numbers.append(row_number)
            stats.data_rows_read += 1

            if len(row_numbers) >= batch_size:
                stats.batches += 1
                yield _build_batch(columns, sheet_name, row_numbers)
                logger.info("%s: %d rows read", sheet_name, stats.data_rows_read)
                columns = [[] for _ in _SOURCE_COLUMNS]
                row_numbers = []

        if row_numbers:
            stats.batches += 1
            yield _build_batch(columns, sheet_name, row_numbers)
        logger.info(
            "%s: finished, %d data rows read, %d blank rows skipped",
            sheet_name,
            stats.data_rows_read,
            stats.blank_rows_skipped,
        )
    finally:
        workbook.close()
